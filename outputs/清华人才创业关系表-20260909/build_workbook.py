from pathlib import Path
import re, json, math, hashlib
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
REF = ROOT.parent / '高校信息收集/清华高校老师信息.xlsx'
def clean(s):
    s = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', str(s or ''))
    s = re.sub(r'[*`]', '', s)
    s = re.sub(r'(Semantic Scholar检索|DBLP检索|Google Scholar检索|Google Scholar|清华学者库)', '', s)
    return re.sub(r'[ \t]+', ' ', s).strip(' |\n；')
def meta(t,k):
    h=t.split('---')[1] if t.startswith('---') else ''
    m=re.search(r'^'+re.escape(k)+r':\s*([^\n]*)',h,re.M)
    if not m:return ''
    v=m.group(1).strip()
    if not v:
        m=re.search(r'^'+re.escape(k)+r':\s*\n((?:\s*- [^\n]*\n?)+)',h,re.M)
        return '；'.join(re.findall(r'- (.*)',m.group(1))) if m else ''
    return v.strip('[]').replace('"','').replace("'",'')
def section(t,heading):
    m=re.search(r'^## '+heading+r'[^\n]*\n(.*?)(?=^## |\Z)',t,re.M|re.S)
    return m.group(1).strip() if m else ''
def table_rows(t):
    return [[clean(c) for c in l.strip().strip('|').split('|')] for l in t.splitlines() if l.startswith('|') and not re.match(r'^\|[\s:|-]+$',l)]
def field(t,*keys):
    for row in table_rows(t):
        if len(row)>1 and row[0] in keys:return row[1]
    return ''
def summary(t):
    s=section(t,'一句话判断')
    return clean(s.split('\n\n')[0])
def urls(t):
    return list(dict.fromkeys(re.findall(r'https?://[^\s)）<>；"|]+',t)))
people={}
for p in sorted((ROOT/'wiki/people').glob('*.md')):
    t=p.read_text(); n=meta(t,'name') or p.stem.split('_')[0]
    key='张涛（EIR院长）' if '张涛_清华具身' in p.name else n
    people[key]={'name':key,'path':p,'text':t,'status':meta(t,'status'),'tier':meta(t,'watch_tier'),
       'aff':meta(t,'affiliations'),'domains':meta(t,'domains'),'role':meta(t,'role_type'),
       'confidence':meta(t,'confidence'),'date':meta(t,'last_verified'),'companies':meta(t,'related_companies'),
       'identity':field(t,'当前身份','身份及公司','公开身份','高校身份'),
       'background':field(t,'清华关联','教育/任职关联','高校关联','学历'),
       'mentor':field(t,'导师/实验室','导师/单位','导师/团队','导师','校内导师','项目指导','实验室'),
       'direction':field(t,'方向','研究方向'),'summary':summary(t)}
companies=[]
for p in sorted((ROOT/'wiki/companies').glob('*.md')):
    t=p.read_text(); companies.append({'name':meta(t,'name') or p.stem,'path':p,'text':t,
      'business':field(t,'主营业务','主营方向','业务方向'),'summary':summary(t)})
def company_lookup(name):
    name=name.strip()
    hits=[c for c in companies if name==c['name'] or name==c['path'].stem]
    if not hits:hits=[c for c in companies if name and (name in c['name'] or name in c['path'].stem)]
    return hits[0] if hits else None
original=list(load_workbook(REF).active.values)
original_names={('姜峣' if r[0]=='姜岣' else r[0]):list(r) for r in original[1:]}
exclude={
 '李弘扬':'库内一手身份为香港大学；清华外聘未获公开支持。',
 '刘年丰':'本人中科大/中科院路径；中科第五纪通过孙富春保留清华入口。',
 '张翼':'库内记为上海交通大学本硕；未获清华关系证据。',
 '祝毅晨':'库内学历为多伦多大学；清华关系未获支持。',
 '李银川':'库内为北理工本博、哥大联培；清华关系未获支持。',
 '姜旭':'个人清华背景未获支持。',
 '宋知珩':'个人清华计算机博士/MBA均待核；公司清华来源不替代个人学籍。',
 '武伟':'公司团队与FIB-Lab关联，不足以确认本人清华身份或论文作者归属。',
 '张果':'清华本科、MIT博士及产品线均缺公开支持。',
 '闵宇恒':'已具名确认创始人；团队源自SIGS，但个人清华学籍/聘任待核。',
 '冯瑶':'仍为拟入职且实际入职未核；本次不计已具清华身份人物。',
 '陈威廉':'转赴清华的报道不等于已确认的学籍/聘任。',
 '廖浩然':'Incoming PhD，实际入学未核；不计在校学生。',
}
# Accept direct study, employment, postdoctoral or officially documented cultivation ties.
extra_accept=set('李健雄 李晓飞 周博文 喻超 杨思成 孙杰 高培中 杨洪兵 王磊 姚卯青 李一帆 都大龙 徐志根 闵伟 罗瑞琨 董浩 毛书翰 邵天兰 唐文斌'.split())
accepted={n for n,p in people.items() if ('清华' in p['aff'] or n in extra_accept) and n not in exclude}
# The original sheet also contains a named, early student signal without a person page.
mp=ROOT/'wiki/maps/清华大学人工智能学院高潜观察池.md'
people['李昊展']={'name':'李昊展','path':mp,'text':mp.read_text(),'status':'student','tier':'T3','aff':'清华大学人工智能学院',
 'domains':'embodied-ai, reinforcement-learning','role':'student','confidence':'single_source','date':'2026-09-08','companies':'',
 'identity':'人工智能学院2025级博士（原始名单口径）','background':'未央书院双学士（数理+机械），原始名单口径',
 'mentor':'李一鸣（原始名单口径，正式培养材料待核）','direction':'VLA、大模型强化学习（PRM）','summary':'原始名单中的早期学生线索；导师李一鸣已创办厘清智能，个人论文与学籍待补证。'}
accepted.add('李昊展')
# Explicit corrections take precedence over stale YAML and the reference sheet.
overrides={
 '詹仙园':dict(identity='AIR研究人员；本溯智能共同创业者（具体职称/公司职务待核）',companies='本溯智能',mentor='与李健雄长期科研合作并共同创业；李健雄自述博士导师为张亚勤，不把詹仙园写成正式学位导师',direction='机器人端侧自适应、In-Context Learning与决策模型'),
 '李健雄':dict(identity='本溯智能创办者；专访自述已博士毕业，毕业时间待核',background='清华AIR博士研究经历；本科为西安交通大学；毕业及学位授予细节待校方补证',mentor='自述师从张亚勤；与詹仙园长期科研合作并共同创业，两种关系分开',companies='本溯智能'),
 '刘淼':dict(identity='人工智能学院PI；未公开创业线索待核',companies='未公开公司名（原始名单线索）'),
 '宋悦':dict(companies='昊瞳科技（仅原始名单，公开主体未核）'),
 '王同翰':dict(companies='昊瞳科技（仅原始名单，公开主体未核）',background='清华人工智能学院PI；原表“姚班硕士”不予确认为学历',mentor='Delta-I Lab；与宋悦共同创业仅为原始名单线索，非师承'),
 '丁宁':dict(identity='电子工程系助理教授（原始名单，官网待核）；创业身份待核',companies='自然意志（仅原始名单，注册及创始关系待核）'),
 '陈勇超':dict(identity='人工智能学院助理教授；公司关系待核',companies='超衍智能（关系仅原始名单支持）'),
 '姜峣':dict(identity='机械工程系制造工程研究所副研究员',companies='橡木果机器人（关联角色及同一人核验待补）'),
 '梁斌':dict(companies='弦引未来（技术源头，非已确认创始人）',direction='绳驱柔性机械臂、空间机器人与遥操作',mentor='导航与控制研究所负责人；王学谦公开称其为导师；郭冠求是王学谦学生'),
 '孙富春':dict(companies='中科第五纪（首席科学家，二手来源）'),
 '李升波':dict(identity='清华车辆与运载学院长聘教授；光象科技联合创始人/首席科学家（媒体）',aff='清华大学车辆与运载学院'),
 '张涛':dict(identity='光象科技创始人兼CEO',background='清华车辆工程博士；勿与EIR院长张涛混淆',aff='清华大学车辆与运载学院',mentor='与李升波为车辆学院师兄弟（媒体口径）；共同导师姓名未核，不写师生'),
 '崔森':dict(background='钱学森力学班本科；清华水木学者/助理研究员口径',mentor='郑泉水为钱班创办首席教授；原表“钱班博士、博士师从郑泉水”未获确认',direction='自进化世界模型',companies='无界矩阵'),
 '胡英东':dict(companies='千寻智能（前沿技术探索任职，非创始人）'),
 '佴瑞乾':dict(companies='未见本人公开创业；高阳组博士生'),
 '谭恒楷':dict(companies='Motus（研究项目，不是本人创业公司）',direction='具身智能世界模型；公司任职未核'),
 '许华哲':dict(companies='破壳机器人（现创办）；星海图（历史联合创始/前首席科学家）'),
 '姚颂':dict(companies='深鉴科技（历史）；东方空间（历史）；正行创新（关联线索待核）'),
 '于超':dict(companies='正行创新（原始名单关联，确切角色待核）'),
 '兰艳艳':dict(companies='艾斐智药AIPher（关联创办；具体职务待核）'),
 '周谷越':dict(companies='求之科技AIRBOT（AIR合作/技术来源；个人职务待核）',mentor='AIR机器人方向、DISCOVER网络；本人学位导师待核'),
 '姚远':dict(companies='面壁智能（多模态首席科学家，非联合创始人）'),
 '徐葳':dict(companies='华控清交；慧安金科（均为首席科学家/转化关联）'),
 '姚期智':dict(companies='小马智行（顾问）；华控清交（成果转化）；南京图灵研究院（平台发起）'),
 '汪玉':dict(companies='深鉴科技（联合创始人）；无问芯穹（发起人，非官网列名首席科学家）',mentor='汪玉课题组/NICS-EFC；姚颂本科指导、单羿/夏立雪/戴国浩博士培养；于超师承仅原表口径'),
 '王学谦':dict(companies='知有无界（学生郭冠求创业，本人创始身份未核）'),
 '苏航':dict(companies='自变量机器人（产学研专项首席科学家，不等于公司职务）'),
 '赵明国':dict(companies='加速进化（首席科学家/实验室来源，非据此确认创始人）'),
 '陈峰':dict(companies='千诀科技（首席技术顾问）'),
 '刘烨斌':dict(companies='影身智能（同窗技术合作）；海信视像（科研合作）'),
 '黄高':dict(companies='小白世纪（库内记联合创始人兼首席科学家；起止时间待核）'),
 '李晓飞':dict(identity='深朴智能创始人兼CEO；曾参与创办智行者',background='清华汽车工程系本科、博士（库内报告）',aff='清华大学车辆与运载学院',companies='深朴智能；智行者（历史）',mentor='具体导师与清华官方学位材料待补'),
 '李一鸣':dict(mentor='NYU博士导师Chen Feng；李昊展为名单所记学生，正式培养关系待核；勿与弋力Li Yi混淆'),
 '刘潇':dict(companies='智谱AI（论文双署名/研发关联，非据此确认创始人）'),
 '董胤蓬':dict(companies='未见本人公开创业；朱军博士生，T-STAR Lab PI'),
 '张金涛':dict(companies='SageAttention（开源研究项目，非公司）；生数等采用其成果'),
 '唐文斌':dict(background='原力灵机/相关报道记清华校友，具体院系学位待官方补证',mentor='范浩强早期信息学教练及职业引荐者；非学位导师'),
 '董浩':dict(background='前清华火神机器人足球队队长；加入加速进化，具体学位待核',mentor='火神队/赵明国指导网络；不据此确认学位导师'),
 '李卓然':dict(identity='清华IIIS博士后',mentor='博士及博士后指导黄隆波（Longbo Huang），本人主页口径'),
 '吕雪广':dict(identity='清华IIIS博士后',mentor='博士导师Christopher Amato（Northeastern）；清华博士后与高阳合作，非高阳博士生'),
 '陈乐偲':dict(identity='清华IIIS博士生，2023级（IIIS官方）',mentor='博士导师张景昭；AI Theory Group；此前复旦本科由Luo Luo指导'),
 '张亚勤':dict(companies='本溯智能（学生李健雄创业）；百度/微软（历史高管任职）',mentor='AIR创始院长；李健雄在专访自述师从张亚勤攻读博士；非据此确认公司职务'),
 '王鹤':dict(background='清华微电子与纳电子学系2010级本科校友（校方补证）；现北大教师',aff='清华大学微电子与纳电子学系（历史学籍）'),
 '刘子鸣':dict(mentor='LIU Lab PI；个人学位导师库内未明确'),
 '刘淼':dict(identity='人工智能学院PI；未公开创业线索待核',companies='未公开公司名（原始名单线索）',mentor='MEOW LAB PI；与冯瑶共同创业仅原始名单线索，后者清华入职待核'),
 '宋悦':dict(companies='昊瞳科技（仅原始名单，公开主体未核）',mentor='结构表征学习方向；与王同翰共同创业仅原始名单线索，非师承'),
 '唐杰':dict(mentor='KEG/知识工程实验室；刘潇博士导师；与李涓子、张鹏属智谱技术/创业网络，不推断张鹏博士师承'),
 '刘知远':dict(mentor='THUNLP；博士导师孙茂松（库内口径仍待学位材料补证）；曾国洋本科导师'),
 '孙茂松':dict(mentor='THUNLP创建者/长期负责人；刘知远博士培养关系按库内待补证口径'),
 '姚远':dict(companies='面壁智能（多模态首席科学家，非联合创始人）',mentor='人工智能学院多模态智能课题组；面壁多模态首席科学家，与刘知远属公司科研网络'),
 '张卫强':dict(mentor='SATLab；EIR数据与算力研究中心主任；不推断与同系创业者的师承'),
 '封硕':dict(identity='自动化系副教授（库内记录）；公司具体角色待核',companies='幂级智能Dense AI（原始名单，创始/公司关系待核）'),
}
for n,vals in overrides.items():people[n].update(vals)
accepted.add('张涛')
# A broad interview question does not independently verify Tang's own Tsinghua degree.
accepted.discard('唐文斌')

dept_order=['交叉信息研究院','人工智能学院','AIR','计算机系','电子工程系','微纳电子系（历史）','自动化系','深圳国际研究生院','机械工程系','车辆与运载学院','软件学院','精密仪器系','生物医学工程学院','工程物理系','物理系','航天航空学院','院系待核']
def dept(p):
    unknown={'杨思成','孙杰','王磊','李一帆','都大龙','徐志根','闵伟','周博文','杨洪兵','姚卯青','喻超','唐文斌'}
    if p['name'] in unknown:return '院系待核'
    if p['name']=='王鹤':return '微纳电子系（历史）'
    s=p['aff']+'；'+p['background']
    keys=[('交叉信息','交叉信息研究院'),('人工智能学院','人工智能学院'),('智能产业','AIR'),('AIR','AIR'),('计算机','计算机系'),('电子工程','电子工程系'),('自动化','自动化系'),('深圳','深圳国际研究生院'),('机械','机械工程系'),('车辆','车辆与运载学院'),('汽车','车辆与运载学院'),('软件','软件学院'),('精密','精密仪器系'),('精仪','精密仪器系'),('生物医学','生物医学工程学院'),('神经工程','生物医学工程学院'),('工程物理','工程物理系'),('物理系','物理系'),('航天','航天航空学院')]
    # First affiliation is the current/primary anchor; history remains in relationship text.
    first=re.split(r'[,；]',p['aff'])[0]
    for a,b in keys:
        if a in first:return b
    for a,b in keys:
        if a in s:return b
    return '院系待核'
groups=[
 ('01 IIIS｜陈建宇·ISR师生创业链','陈建宇 胡钰承 姜哲源'),
 ('02 IIIS｜高阳·EVAR培养与千寻','高阳 佴瑞乾 胡英东 吕雪广'),
 ('03 IIIS｜星海图/破壳创业网络','许华哲 赵行 高继扬'),
 ('04 IIIS/计算机｜弋力·兴军亮与银河研究网络','弋力 张智楷 薛晗 兴军亮 Haofei Lu 王鹤'),
 ('05 IIIS｜姚期智·小马与转化平台','姚期智 楼天城 彭军 徐葳'),
 ('06 IIIS｜李建培养与理论/系统','李建 吕凯风 李卓然 张景昭 陈乐偲 张焕晨 高鸣宇 杜韬 徐梦迪'),
 ('07 姚班｜原力灵机竞赛与职业网络','唐文斌 范浩强'),
 ('08 AI学院｜李一鸣培养链','李一鸣 李昊展'),
 ('09 AI学院｜青年PI（同院为主）','邹雪妍 刘淼 宋悦 王同翰 刘子鸣 郭钰铎 李佳 陈勇超'),
 ('10 AIR｜张亚勤培养·詹仙园合作·本溯','张亚勤 詹仙园 李健雄'),
 ('11 AIR｜机器人与产业转化','周谷越 陈亦伦 赵昊 马维英 刘洋（AIR）'),
 ('12 AIR｜AI制药与科学智能','兰艳艳 聂再清 周浩'),
 ('13 计算机｜朱军·TSAIL培养链','朱军 董胤蓬 张金涛 谭恒楷 苏航'),
 ('14 计算机｜THUNLP师承/公司网络','孙茂松 刘知远 曾国洋 姚远'),
 ('15 计算机｜KEG·智谱网络','唐杰 李涓子 张鹏 刘潇'),
 ('16 计算机/EIR｜机器人与智能系统','孙富春 刘华平 张钹 胡晓林 邓志东 刘永进 崔鹏'),
 ('17 电子/SIGS｜汪玉·NICS-EFC师生网络','汪玉 姚颂 单羿 夏立雪 戴国浩 于超 韩松'),
 ('18 电子｜贝塔无限同团队','刘武龙 陶帅'),
 ('19 电子｜李勇·ChaosNexus项目指导','刘畅 赵博浩'),
 ('20 电子｜视觉感知（同院/方向）','方璐 王贵锦 王生进 王陈玉珩 刘芳甫 张卫强 丁宁 王潜 邵天兰 李力耘'),
 ('21 自动化/SIGS｜梁斌→王学谦→郭冠求','梁斌 王学谦 郭冠求'),
 ('22 自动化｜赵明国·火神队与加速进化','赵明国 程昊 董浩'),
 ('23 自动化｜陈峰·千诀技术团队','陈峰 高海川'),
 ('24 自动化｜莫一林·金戈创业伙伴','莫一林 金戈'),
 ('25 自动化｜同窗/产业合作与影身','刘烨斌 闵伟 季向阳'),
 ('26 自动化/计算机｜德塔共同创业网络','黄思远 马晓健'),
 ('27 自动化/EIR｜控制与视觉（同院）','张涛（EIR院长） 何潇 鲁继文 朱松纯 黄冠 封硕 黄高 许闻达'),
 ('28 SIGS/航院｜郑泉水·钱班培养网络','郑泉水 崔森'),
 ('29 SIGS｜触觉与创新项目','丁文伯 高培中'),
 ('30 机械｜任晓雨·毛书翰同寝室创业','任晓雨 毛书翰'),
 ('31 机械｜制造工程所','姜峣 李曙光 唐晓强'),
 ('32 机械｜机械电子所','徐静 杨东超 吴丹'),
 ('33 机械｜设计工程所','赵慧婵 刘辛军 谢福贵 季林红'),
 ('34 车辆｜光象与汽车技术创业','李升波 张涛 董汉 李晓飞 秦深涛'),
 ('35 软件｜龙明盛·吴海旭','龙明盛 吴海旭'),
 ('36 精仪｜CBICR·晰见','杨哲宇 王韬毅'),
 ('37 生医｜洪波指导·灵犀','闫宇翔'),
 ('38 工物｜SUNIST与星环聚能','谭熠 陈锐 高喆 蒲以康 王侃'),
 ('39 物理｜量子创业','胡晓晓'),
]
groupmap={}
for g,names in groups:
    ns=names.split()
    if 'Haofei Lu' in names:ns=[n for n in ns if n not in ['Haofei','Lu']]+['Haofei Lu']
    for i,n in enumerate(ns):groupmap[n]=(g,i)
def sortkey(n):
    if n in groupmap:return (*groupmap[n],n)
    d=dept(people[n]);return ('90 '+str(dept_order.index(d)).zfill(2)+' '+d,0,n)
def group(n):return sortkey(n)[0]
def relationship(n):
    p=people[n]
    bits=[group(n), '身份：'+(p['background'] or p['identity'] or p['aff']), '师承/团队：'+(p['mentor'] or '库内未明确导师或实验室')]
    return '\n'.join(bits)

uncertain=set('刘淼 宋悦 王同翰 丁宁 陈勇超 姜峣 于超 周谷越 封硕'.split())
public_founders=set('陈建宇 胡钰承 姜哲源 高阳 许华哲 赵行 高继扬 楼天城 彭军 范浩强 唐文斌 李一鸣 詹仙园 李健雄 陈亦伦 兰艳艳 聂再清 朱军 刘知远 曾国洋 唐杰 李涓子 张鹏 汪玉 姚颂 单羿 夏立雪 戴国浩 韩松 刘武龙 陶帅 王陈玉珩 王潜 邵天兰 郭冠求 程昊 高海川 莫一林 金戈 闵伟 黄思远 马晓健 黄冠 封硕 黄高 崔森 丁文伯 高培中 任晓雨 毛书翰 李升波 张涛 董汉 李晓飞 秦深涛 杨哲宇 王韬毅 闫宇翔 谭熠 陈锐 胡晓晓 王栋 王冠 喻超 杨思成 孙杰 杨洪兵 王磊 姚卯青 李一帆 都大龙 周博文'.split())
public_founders.discard('封硕');public_founders.add('王鹤')
# Main preserves the input population where in scope, adds company-linked people and useful relational anchors.
mainset=(set(original_names)&accepted)|{n for n in accepted if people[n]['companies'] and n not in {'马维英'}}|public_founders
mainset&=accepted
mainset|=set('陈建宇 高阳 许华哲 赵行 弋力 兴军亮 姚期智 李建 朱军 孙茂松 唐杰 李涓子 汪玉 赵明国 王学谦 陈峰 郑泉水 龙明盛 张钹'.split())&accepted
annexset={n for n in accepted if n not in public_founders and (people[n]['tier'] in ['T1','T2'] or n in original_names)}
annexset-=set('姚期智'.split()) # Already a disclosed platform founder, not a pre-venture prospect.
main_names=sorted(mainset,key=sortkey); annex_names=sorted(annexset,key=sortkey)

def linked_companies(p):
    paths=[]
    for rel in re.findall(r'\]\(([^)]*companies/[^)]*\.md)\)',p['text']):
        f=(p['path'].parent/rel).resolve()
        if f.exists():paths.append(f)
    # Include confirmed overrides which may not be in stale YAML.
    for c in companies:
        short=c['name'].split('（')[0].split('(')[0]
        if len(short)>2 and short in p['companies']:paths.append(c['path'])
    return list(dict.fromkeys(paths))
def evidence(p):
    paths=[p['path']]+linked_companies(p)
    refs=[]
    for path in paths:
        t=path.read_text()
        raw=[]
        for rel in re.findall(r'\]\(([^)]*raw/[^)]*\.md)\)',t):
            rp=(path.parent/rel).resolve()
            if rp.exists():raw.append(rp)
        refs.extend(raw)
    return paths,list(dict.fromkeys(refs))
def comment(p,extra=''):
    ps,raws=evidence(p)
    return Comment('资料截至库内核验日期：'+p['date']+'\n'+extra+'\n知识库：\n'+'\n'.join(str(x) for x in ps)+'\n原始来源：\n'+'\n'.join(str(x) for x in raws)+'\n公开链接：\n'+'\n'.join(urls(p['text'])[:6]),'知识库整理')
def venture_state(n):
    if n in uncertain:return '未公开确认：内部/关联线索待核'
    if n in public_founders:return '库内已有创业/创始团队记录；具体职务与时效见来源'
    p=people[n]
    if p['companies']:return '已有产业、任职或研究关联；未据此确认本人独立创业'
    return '本库未记录本人公开创业；不等于现实中没有创业'
def business(p):
    out=[]
    for f in linked_companies(p):
        c=next((c for c in companies if c['path']==f),None)
        if c and c['business']:out.append(c['name']+'：'+c['business'])
    return '\n'.join(out[:3])
def signals(p):
    sec=section(p['text'],'关键信号') or section(p['text'],'可跟踪信号') or section(p['text'],'身份与独立信号')
    good=[]
    for r in table_rows(sec):
        if len(r)<2 or r[0] in ['类型','信号','独立信号','建页动作','任职','研究方向','清华职务','组织身位','实验室','作者证据','开源入口']:continue
        good.append(r[0]+'：'+r[1])
    return '\n'.join(good[:2]) or p['summary'] or p['direction'] or '库内仅有早期线索，独立成果待补'
manual_signal={
 '刘子鸣':'LIU Lab科学智能/物理与AI交叉；KAN架构线索（原始名单与学院方向记录）。',
 '刘淼':'MEOW LAB；第一视角、多模态、人本智能研究；共同创业仅原始名单，未公开。',
 '宋悦':'结构化表征学习、科学智能；模型压缩/ASVD为名单中的潜在转化判断。',
 '王同翰':'Delta-I Lab，多智能体/机制设计/Human-AI经济学；“昊瞳科技”公开主体未核。',
 '周浩':'DAPO强化学习系统线索（原始名单）；与具体公司的创始/任职关系未核。',
 '张金涛':'SageAttention与SageAttention2核心作者；首届张钹奖学金；已有模型/工程采用记录。',
 '刘潇':'P-tuning、GLM-130B、AgentBench第一作者；首届张钹奖学金；智谱双署名。',
 '董胤蓬':'AI安全对抗方法与ARES/MultiTrust/T2VSafetyBench；朱军博士培养链。',
 '佴瑞乾':'OneTwoVLA与HuMI共同一作；Amazon FAR人形机器人操作实习。',
 '胡英东':'OneTwoVLA、Data Scaling Laws共同一作；本人主页已披露博士毕业和千寻任职。',
 '吕雪广':'JAIR/AAAI多智能体强化学习论文；清华IIIS博士后，与高阳合作。',
 '高喆':'SUNIST/磁约束核聚变平台；与谭熠同平台，库内未记录本人公司化。',
 '蒲以康':'工物系等离子体/聚变研究网络；库内个人产业转化证据薄，需补具体成果与项目角色。',
 '王侃':'工物系/REAL Lab入口；核能与先进能源方向，具体创业转化证据待补。',
}
direct_comparators={
 '佴瑞乾':[('高阳','导师；高阳已参与创办千寻智能，最直接的培养链样本')],
 '胡英东':[('高阳','博士导师；本人已进入千寻，适合产业合作与人才网络接触')],
 '吕雪广':[('高阳','当前博士后合作教师；不是其博士导师')],
 '张金涛':[('朱军','导师；生数科技可作为同一培养链转化参照'),('夏立雪','算力/推理效率方向相近；不同院系，无师承推断')],
 '谭恒楷':[('朱军','导师与TSAIL；生数科技为近关系参照，Motus署名不证明任职')],
 '董胤蓬':[('朱军','博士导师；同一培养链的基础模型创业参照')],
 '刘潇':[('唐杰','导师/KEG；智谱为直接实验室转化参照')],
 '孙茂松':[('刘知远','培养/THUNLP网络；面壁为实验室转化案例，具体师承以人物页边界为准')],
 '姚远':[('刘知远','面壁共同科研/企业网络；姚远是多模态首席科学家，非联合创始人')],
 '弋力':[('王鹤','同为Guibas博士培养网络，且已有银河研究合作；本人非据此认定创业')],
 '张智楷':[('王鹤','通过导师弋力和银河研究合作相连；不是王鹤博士生')],
 '薛晗':[('王鹤','通过导师弋力和银河研究合作相连；学籍为人工智能学院')],
 'Haofei Lu':[('王鹤','银河研究网络；校内导师兴军亮、企业研究指导弋力，角色分开')],
 '兴军亮':[('王鹤','银河合作研究网络；不等同联合创业或师承')],
 '李昊展':[('李一鸣','原始名单中的导师；厘清为近关系参照，学生创业与加入公司未核')],
 '王学谦':[('郭冠求','本人学生已创办知有无界；梁斌→王学谦→郭冠求链')],
 '梁斌':[('郭冠求','学生王学谦的学生；知有无界已创业，弦引未来另属本人技术源头转化')],
 '赵明国':[('程昊','本科实验室/火神队培养与加速进化产业网络')],
 '陈峰':[('高海川','曾任其类脑双臂课题组组长；本人为千诀顾问，不推断博士师生')],
 '刘烨斌':[('闵伟','清华同窗及影身智能技术合作；本人股权/公司职务待核')],
 '郑泉水':[('崔森','钱班创办教授与本科培养网络；不确认为博士师承')],
 '吴海旭':[('李一鸣','世界模型/学习算法转化方向相近，跨院系；不是师承'),('黄冠','时序/物理预测与世界模型相近；不同培养链')],
 '高喆':[('谭熠','同SUNIST平台，星环聚能为明确的同平台创业样本')],
 '李佳':[('刘知远','大模型工具/产品化路径相近，跨院系；非十联合产品不等于本人创业')],
 '高鸣宇':[('夏立雪','AI体系结构/算力基础设施画像相近；本人IIIS、夏为电子系，不认定同门')],
 '徐梦迪':[('高阳','同IIIS且机器人学习/适应方向相近；仅同院与技术路径参照')],
 '杜韬':[('陈建宇','同IIIS，机器人仿真/设计与本体创业相邻；非同门证据')],
}
priority_direct=set(direct_comparators)
direct_comparators.update({
 '于超':[('姚颂','原始名单称同出汪玉门下并共同发起正行创新；姚颂历史创业已公开，于超/正行角色仍待核')],
 '张亚勤':[('李健雄','本人专访自述师从张亚勤，并与詹仙园共同创办本溯；不据此推断张亚勤持股/顾问')],
 '苏航':[('王潜','已有自变量产学研专项合作；苏航专项首席科学家不等于公司创始人')],
 '刘子鸣':[('兰艳艳','科学智能→AI制药为跨院转化路径参照，具体技术不同，无师承证据'),('李一鸣','同AI学院青年PI已创业；仅组织与职业阶段参照')],
 '郭钰铎':[('兰艳艳','AI4Science转化方向可参照；非同院或同门'),('李一鸣','同AI学院青年PI创业，技术路线不等同')],
 '方璐':[('王陈玉珩','同电子系、光学/视触觉感知画像相近；未确认导师关系')],
 '王贵锦':[('王陈玉珩','同电子系，视觉计算与机器人感知转化相近；未确认师承')],
 '王生进':[('王潜','同电子系、视觉与具身智能方向相近；仅同院/方向参照')],
 '张卫强':[('王陈玉珩','同电子系、感知与数据基础设施为相邻方向；不代表音频技术相同')],
 '刘畅':[('高继扬','同电子系、物理世界建模与机器人方向相邻；本人项目指导李勇，非同门证据')],
 '赵博浩':[('高继扬','同电子系、物理世界建模与机器人方向相邻；本人项目指导李勇，非同门证据')],
 '周谷越':[('陈亦伦','同AIR机器人方向，学术/产业交叉经历相近；具体师承未确认'),('詹仙园','同AIR模型转化样本；不等同共同创业')],
 '赵昊':[('陈亦伦','同AIR，机器人/自动驾驶技术画像相邻；未核直接培养关系')],
 '刘华平':[('马晓健','同计算机系，机器人感知/学习到具身公司的路径参照；非师承')],
 '孙富春':[('马晓健','同计算机系、具身方向参照；本人的中科第五纪首席科学家身份另列，不当成已创办')],
 '崔鹏':[('朱军','同计算机系，基础方法到模型公司转化可参照；稳定学习与生成模型路线不同')],
 '季向阳':[('黄冠','同自动化系、视觉/感知画像相邻；未确认师承')],
 '朱松纯':[('黄思远','BIGAI产业研究网络、德塔联合创业记录；不因UCLA/BIGAI推定博士师承')],
 '张涛（EIR院长）':[('程昊','同自动化系机器人创业样本；与光象CEO张涛不是同一人')],
 '何潇':[('程昊','同自动化系、控制与机器人技术转化参照；未确认师承')],
 '鲁继文':[('黄冠','同自动化系、机器视觉/具身方向相近；学位导师与实验室关系不可推定')],
 '徐葳':[('楼天城','同IIIS的另一产业化路径：小马智行；本人已任华控清交/慧安首席科学家，非待产业化新人')],
 '李建':[('楼天城','同IIIS创业样本，非据此认定指导关系；李建组另有金一飞→Xtech的学生去向，角色/主体待补')],
 '张焕晨':[('夏立雪','AI数据系统与异构算力基础设施属相邻底层方向；跨院、无师承证据')],
 '龙明盛':[('李一鸣','时序预测/物理建模与世界模型方向相邻；跨院且技术产品不同')],
 '封硕':[('黄冠','同自动化系且世界模型方向相邻；黄冠创业已有媒体，封硕的幂级智能关系仍待核')],
})
for n in ['姜峣','李曙光','唐晓强','徐静','吴丹','刘辛军','谢福贵','季林红']:
    direct_comparators[n]=[('任晓雨','同机械系，机器人技术转化参照；不推断为本人学生或同一课题组')]
direct_comparators['杨东超']=[('任晓雨','同机械系，机器狗/传感器产业化与机器人本体方向相邻，非师承'),('王陈玉珩','触觉传感器转化方向相近；跨院比较')]
direct_comparators['赵慧婵']=[('任晓雨','同机械系机器人创业样本；柔性机器人与人形本体路线不同')]
def comparisons(n):
    if n in direct_comparators:return direct_comparators[n]
    p=people[n]; ds=set(p['domains'].split(', ')); candidates=[]
    for cn in public_founders&accepted:
        c=people[cn]; cs=set(c['domains'].split(', ')); same=dept(p)==dept(c)
        overlap=ds&cs-{'ai','robotics','frontier-talent'}
        score=(4 if same else 0)+len(overlap)
        if score:candidates.append((score,cn,same,overlap))
    candidates.sort(reverse=True)
    ans=[]
    for score,cn,same,overlap in candidates[:2]:
        basis='同'+dept(p)+'；仅院系参照，未确认师承' if same else '研究方向相近；跨院系，未确认师承或合作'
        ans.append((cn,basis))
    return ans
def priority(n):
    p=people[n]
    if n in uncertain or p['tier']=='T3' or p['confidence']=='single_source':return 'P3 先补证'
    if n in priority_direct or p['tier']=='T1':return 'P1 优先了解'
    return 'P2 定向了解'
def next_action(n):
    if n in uncertain:return '先确认主体名称、本人角色与公开披露；未核前不按创始人安排接触。'
    if n=='胡英东':return '围绕千寻现岗位与技术路线接触；区分产业合作、人才交流和独立创业意愿。'
    if priority(n).startswith('P3'):return '先补官方履历、具名成果和导师材料，再判断接触优先级。'
    if n in direct_comparators:return '核实当前身份与项目投入；从已记录导师/团队链了解转化方向和接触入口。'
    return '围绕代表成果了解应用痛点、学生团队与产业化意愿；确认公司角色。'

wb=Workbook();main=wb.active;main.title='主表_清华人物与公司';annex=wb.create_sheet('附表_未公开创业重点人物');info=wb.create_sheet('编制说明');sources=wb.create_sheet('来源索引')
main.append(list(original[0])+['清华身份与师承/团队关系'])
for n in main_names:
    p=people[n];orig=original_names.get(n);d=dept(p)
    # Keep the schema and two original department semantics; correct factual content from the corpus.
    direction=p['direction'] or (orig[5] if orig else '') or p['domains'].replace(', ',' / ')
    notes=business(p) or signals(p)
    second=orig[6] if orig else ''
    if second:second=str(second)+'（原表分类，非已核技术结论）'
    if n=='崔森':second='未按WAM归类；官方口径为自进化世界模型'
    identity=p['identity'] or p['summary'] or '具体职务待核'
    role='学术' if p['role'] in ['professor','student','researcher'] or '教授' in identity else '工业'
    row=[n,p['aff'] or '清华背景；具体院系待核',d,identity,role,direction,second,venture_state(n),notes,p['companies'] or '本库未记录公开创业项目',relationship(n)]
    main.append(row)
    main.cell(main.max_row,1).comment=comment(p)
    main.cell(main.max_row,11).comment=comment(p,'分组是阅读顺序：导师、项目指导、同窗、创业伙伴、同院系分别标注，不等价于同门。')
    if orig:main.cell(main.max_row,8).comment=Comment('参考表原始行：\n'+json.dumps(orig,ensure_ascii=False)+'\n新表依知识库订正；原文件未改动。','参考表核对')
annex.append(['人名','清华院系（主归类）','当前身份/阶段','重点方向','重要性依据（库内信号）','清华身份与师承/团队关系','本人创业/产业状态','可参考创业：相似画像或同院系人物','参照强度与证据边界','建议接触优先级','建议接触切入点/待核项','库内核验日期'])
for n in annex_names:
    p=people[n];pairs=comparisons(n)
    refs='\n'.join(cn+' → '+people[cn]['companies'] for cn,b in pairs) or '库内未找到足够明确的相近创业样本'
    bounds='\n'.join(cn+'：'+b for cn,b in pairs) or '不以方向推测公司或师承'
    state=venture_state(n)
    if p['companies']:state+='\n'+p['companies']
    annex.append([n,dept(p),p['identity'] or p['summary'],p['direction'] or p['domains'].replace(', ',' / '),manual_signal.get(n,signals(p)),relationship(n),state,refs,bounds,priority(n),next_action(n),p['date']])
    annex.cell(annex.max_row,1).comment=comment(p)
    annex.cell(annex.max_row,8).comment=Comment('\n\n'.join(cn+'\n'+str(people[cn]['path'])+'\n'+b for cn,b in pairs),'创业参照依据')

info.append(['项目','口径/说明'])
notes=[
 ('用途','基于参考表与本知识库整理清华人物、公司及培养关系，用于接触优先级研判。未进行新增外部尽调。'),
 ('编制日期','2026-09-09；每人事实时点见来源索引，不把历史材料自动外推为现状。'),
 ('参考文件',str(REF)),
 ('主表结构','原10列名称和顺序全部保留（含两个“学院”列），末尾新增1列。第二列填主组织/归属，第三列规范化归类；内容可按库内证据纠错。'),
 ('主表范围',f'{len(main_names)}位。保留参考表内符合清华范围的人物，补公司关联人物和解释培养链所需导师；主表沿用原表包含未创业人士的设计。'),
 ('附表范围',f'{len(annex_names)}位。以已建T1/T2人物及原表重点观察者为候选，排除明确创业/创始团队与已公开平台发起者；公司任职、顾问、实习、科研合作单独标记。'),
 ('去重',f'两表合计{len(set(main_names)|set(annex_names))}位不同人物；跨表重合{len(set(main_names)&set(annex_names))}位是有意保留的观察视图，不可把两表行数相加当总人数。'),
 ('清华范围','本人具有库内支持的清华学习、任职、博士后或具名培养经历。外校现任教师如有清华学位可纳入；仅与清华人合伙、公司声称清华背景，不自动纳入本人。原表弱来源保留者在状态列降级。'),
 ('未公开创业含义','仅表示当前知识库尚无足够公开创业证据，不证明现实中未创业；原始名单的未公开线索、产品合作、产业任职分别写明。'),
 ('分组逻辑','优先同导师/实验室、师生链、同团队/同窗，再按同院系和方向。分组编号用于两表对照；同院、同平台、同项目都不等于正式师承。'),
 ('关系证据','导师/师生、博士后合作、项目指导、企业实习指导、同学/同窗、创业伙伴、行政负责人分开；学位导师未核就不补造。'),
 ('创业项目列','公司、历史公司、科研项目和学生创业落点均明确关系类型；不把“关联公司”一律改成“本人创办”。'),
 ('世界模型分类','保留原表已有分类时明确标注为原表分析口径；新增人未找到该细分类证据则留空，避免把所有机器人项目强归WAM。'),
 ('优先级（分析建议）','P1：已有较强科研/转化信号或直接导师/团队样本，可优先了解；P2：同院/同方向、研究信号较清楚，可定向了解；P3：单一来源、早期名单或创业身份待核，先补证。不是创业意愿/投资价值评分。'),
 ('创业参照（分析建议）','先选库内明确师承/团队/合作案例，否则选同院系或相近方向的已有创业者；参照列用于比较，不新增事实关系。'),
 ('证据阅读','人物姓名与关系单元格附批注；来源索引提供对应知识文章、原始材料、公开URL和核验日期。链接点击可打开本地材料；外部链接未在本轮重新抓取。'),
 ('现有证据局限','库内部分条目的Sources包含二手检索摘要、raw仅指回参考表，不能当成新一手核验。来源索引列明原始材料，具体职称、公司角色和旧任职需要接触前复核。'),
 ('已处理的重要纠错','姜岣→姜峣（机械系）；李升波→车辆学院；高海川→自动化系且陈峰课题组；崔森钱班为本科、博士师承未核；许华哲现破壳/历史星海图；胡英东毕业并任职千寻。'),
 ('公司覆盖边界','以已明确清华个人入口为中心。流形空间、星忆、零次方等公司层面的清华关系，不能用于补造创始人的学籍；未具名清华个人者不强行进入人物主表。'),
 ('未修改内容','原Excel和wiki/raw/index/log均不改写；本工作簿为独立导出物。'),
]
for r in notes:info.append(r)
info.append(['原表范围调整','下列行从原表清华人物主表排除；仅列核对记录，不列入人才统计。'])
for n,reason in exclude.items():
    if n in original_names:info.append([n,reason])
info.append(['公司保留入口','中科第五纪改由清华教授孙富春承接；原创始人刘年丰不计清华人才。'])
sources.append(['人名','主表行号','附表行号','知识库人物/观察池路径','关联公司页面','Raw原始资料路径','人物页公开来源URL','人物页置信度（非逐项）','库内核验日期','证据提示'])
mr={n:i+2 for i,n in enumerate(main_names)}; ar={n:i+2 for i,n in enumerate(annex_names)}
for n in sorted(set(main_names)|set(annex_names),key=sortkey):
    p=people[n];ps,raws=evidence(p)
    sources.append([n,mr.get(n),ar.get(n),str(p['path']), '\n'.join(str(x) for x in ps[1:]),'\n'.join(str(x) for x in raws), '\n'.join(urls(p['text'])[:12]),p['confidence'],p['date'], '有原始名单/单一来源，关键事实需逐项复核' if p['confidence']=='single_source' or n in uncertain else '按具体事实保留媒体、官方和自述边界；不由页面置信度升级全部结论'])
    sources.cell(sources.max_row,4).hyperlink=p['path'].as_uri()
    if mr.get(n):sources.cell(sources.max_row,2).hyperlink=f"#'{main.title}'!A{mr[n]}"
    if ar.get(n):sources.cell(sources.max_row,3).hyperlink=f"#'{annex.title}'!A{ar[n]}"
sr={sources.cell(i,1).value:i for i in range(2,sources.max_row+1)}
for ws in [main,annex]:
    for i in range(2,ws.max_row+1):ws.cell(i,1).hyperlink=f"#'{sources.title}'!A{sr[ws.cell(i,1).value]}"

navy='17365D';teal='DDEEF0';blue='F0F5FA';gold='FFF2CC';white='FFFFFF';border='C7D5E2'
widths={main.title:[15,28,21,44,10,32,29,35,65,47,80],annex.title:[15,22,42,33,57,73,50,55,62,18,48,16],info.title:[26,120],sources.title:[16,12,12,74,76,84,86,24,17,44]}
for ws in wb:
    ws.sheet_view.showGridLines=False;ws.freeze_panes='D2' if ws in [main,annex] else 'B2'
    ws.auto_filter.ref=ws.dimensions
    for i,w in enumerate(widths[ws.title],1):ws.column_dimensions[get_column_letter(i)].width=w
    for c in ws[1]:c.fill=PatternFill('solid',fgColor=navy);c.font=Font(name='等线',size=11,bold=True,color=white);c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=38
    prev=None
    for r in range(2,ws.max_row+1):
        n=ws.cell(r,1).value;g=group(n) if n in people else None;new=g!=prev
        for c in ws[r]:
            c.font=Font(name='等线',size=11,color='17365D' if c.hyperlink else '243746',underline='single' if c.hyperlink else None)
            c.fill=PatternFill('solid',fgColor=white if r%2==0 else blue)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if new and ws in [main,annex]:c.border=Border(top=Side(style='medium',color='7DA6B5'))
        if ws in [main,annex]:
            ws.cell(r,1).fill=PatternFill('solid',fgColor=teal)
            if (ws==annex and str(ws.cell(r,10).value).startswith('P3')) or n in uncertain:
                ws.cell(r,10 if ws==annex else 8).fill=PatternFill('solid',fgColor=gold)
        lines=max(max(len(str(c.value or '').split('\n')),math.ceil(sum(2 if ord(ch)>127 else 1 for ch in str(c.value or ''))/(widths[ws.title][c.column-1]*0.93))) for c in ws[r])
        ws.row_dimensions[r].height=min(360,max(62,lines*15+12)) if ws!=sources else 95
        prev=g
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1';ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='清华人才与创业关系 | &P / &N'
    ws.sheet_view.zoomScale=75 if ws!=info else 90
main.sheet_properties.tabColor='17365D';annex.sheet_properties.tabColor='16858B';info.sheet_properties.tabColor='D8AA43';sources.sheet_properties.tabColor='8196AA'
dest=OUT/'清华人才与创业关系_知识库补全版_2026-09-09.xlsx'
wb.save(dest)
check=load_workbook(dest)
assert [c.value for c in check.worksheets[0][1]][:10]==list(original[0])
assert len(main_names)==len(set(main_names)) and len(annex_names)==len(set(annex_names))
assert not (set(exclude)&(set(main_names)|set(annex_names)))
assert not (public_founders&set(annex_names))
assert all(p['path'].exists() for n,p in people.items() if n in set(main_names)|set(annex_names))
assert all(c.data_type!='f' for ws in check for row in ws for c in row)
audit={'output':str(dest),'main_people':len(main_names),'annex_people':len(annex_names),'unique_people':len(set(main_names)|set(annex_names)),
 'overlap':len(set(main_names)&set(annex_names)),'original_rows':len(original)-1,'original_retained':len(set(original_names)&set(main_names)),
 'excluded_original':{n:r for n,r in exclude.items() if n in original_names},'priority_counts':dict(Counter(priority(n) for n in annex_names)),
 'main_names':main_names,'annex_names':annex_names,'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest()}
(OUT/'生成核对.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in audit.items() if k not in ['main_names','annex_names']},ensure_ascii=False,indent=2))
